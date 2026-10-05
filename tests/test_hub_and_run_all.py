"""scripts/build_hub.py and scripts/run_all.py: the generated README block, banned phrases,
and the pipeline's control flow (no network, no models, no subprocesses)."""
import argparse
import copy
import importlib.util
import json
from pathlib import Path

import pytest

from zeroops import paths

FACTS = json.loads(paths.FACTS.read_text())


def _load(name: str):
    path = paths.SCRIPTS / f"{name}.py"
    spec = importlib.util.spec_from_file_location(f"{name}_script", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


hub = _load("build_hub")
run_all = _load("run_all")


# ---------------------------------------------------------------- build_hub

def test_readme_results_block_matches_facts():
    assert hub.readme_block(FACTS) in paths.README.read_text(), \
        "README Results block is stale: run scripts/build_hub.py"


def test_banned_phrases_match_whole_phrases_only():
    hub.check_banned("labelled unsafe to automate", "test")
    with pytest.raises(SystemExit):
        hub.check_banned("this one is safe to automate", "test")


def test_generated_copy_has_no_banned_phrases():
    for s in hub.headline(FACTS) + hub.caveats(FACTS):
        hub.check_banned(s, "copy")


def test_headline_numbers_come_from_facts():
    b = FACTS["backends"]["jev_native"]
    text = " ".join(hub.headline(FACTS))
    assert f"{b['action_exact']['k']} of {FACTS['n']}" in text
    assert f"act alone on {b['auto']} of {FACTS['n']}" in text


def test_headline_reports_unsafe_automation_if_it_happens():
    f = copy.deepcopy(FACTS)
    f["backends"]["jev_native"]["unsafe_auto"] = 2
    text = " ".join(hub.headline(f))
    assert "including 2 of the" in text and "none of them" not in text


def test_media_helpers_handle_missing_files(tmp_path):
    assert hub.deck_info(tmp_path / "missing.pptx") is None
    assert hub.video_info(tmp_path / "missing.mp4") is None
    assert hub.fmt_duration(146.53) == "2:26"


# ---------------------------------------------------------------- run_all

class FakeRunner(run_all.Runner):
    """Records commands instead of running them; rc_for(line) decides each exit code."""

    def __init__(self, rc_for):
        super().__init__(dry=False)
        self.calls: list[str] = []
        self.rc_for = rc_for

    def cmd(self, argv, *, cwd=None, env=None):
        line = " ".join(Path(a).name for a in argv)
        self.calls.append(line)
        return self.rc_for(line)


def test_no_backends_is_a_note_not_a_failure():
    r = FakeRunner(lambda line: run_all.NO_BACKENDS)
    r.evaluate("Jev API backends", "--backends", "api")
    assert r.notes and not r.problems


def test_a_failed_evaluation_stops_the_pipeline():
    r = FakeRunner(lambda line: 1)
    with pytest.raises(run_all.StepFailed):
        r.evaluate("Jev API backends", "--backends", "api")


def test_kev_started_here_is_stopped_even_when_its_run_fails(monkeypatch):
    monkeypatch.setattr(run_all, "kev_available", lambda: False)
    monkeypatch.setattr(run_all, "start_kev_and_wait", lambda r: True)
    r = FakeRunner(lambda line: 1 if "--only kev" in line else 0)
    with pytest.raises(run_all.StepFailed):
        run_all.run_local(r, argparse.Namespace(start_kev=True), [])
    assert r.calls[-1].endswith("stop_kev.sh")


def test_kev_that_never_came_up_is_still_stopped(monkeypatch):
    monkeypatch.setattr(run_all, "kev_available", lambda: False)
    monkeypatch.setattr(run_all, "start_kev_and_wait", lambda r: False)
    r = FakeRunner(lambda line: 0)
    run_all.run_local(r, argparse.Namespace(start_kev=True), [])
    assert r.problems and r.calls[-1].endswith("stop_kev.sh")
    assert not any("--only kev" in c for c in r.calls)


def test_a_kev_server_already_running_is_left_alone(monkeypatch):
    monkeypatch.setattr(run_all, "kev_available", lambda: True)
    r = FakeRunner(lambda line: 0)
    run_all.run_local(r, argparse.Namespace(start_kev=True), [])
    assert any("--only kev" in c for c in r.calls)
    assert not any("stop_kev" in c for c in r.calls)


def _deck(path, slides):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for texts in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for i, text in enumerate(texts):
            s.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(4), Inches(1)).text_frame.text = text
    prs.save(path)


def test_deck_info_splits_main_and_appendix(tmp_path):
    p = tmp_path / "d.pptx"
    _deck(p, [["Title"], ["SUMMARY", "x"], ["Appendix", "A1 questions"], ["APPENDIX", "y", "A1"],
              ["APPENDIX", "z", "A2"]])
    assert hub.deck_info(p) | {"notes": 0} == {"total": 5, "main": 2, "appendix": 2, "notes": 0}
    _deck(p, [["Title"], ["APPENDIX", "y", "A1"]])  # no divider slide
    assert (hub.deck_info(p)["main"], hub.deck_info(p)["appendix"]) == (1, 1)


def test_race_info_reads_raw_typed_answers(tmp_path):
    trace = {"incident": {"id": "INC-0001"}, "chat_model": "m", "corpus_version": "v",
             "jev": {"ok": True, "answers": {
                 "failure_category": {"type": "choice", "choice": "app"},
                 "root_cause": {"type": "choice", "choice": "bad_deploy"},
                 "proposed_action": {"type": "choice", "choice": "rollback"}}},
             "chat": {"ok": True, "parsed": {"failure_category": "app", "root_cause": "bad_deploy",
                                             "action": "restart"}}}
    p = tmp_path / "t.json"
    p.write_text(json.dumps(trace))
    info = hub.race_info(p)
    assert info["same"] == ["category", "root cause"] and info["differ"] == ["first action"]
    desc, _, warn = hub.race_copy(info, {"corpus_version": "other"}, 12)
    assert "differed on first action" in desc and "earlier version" in warn
