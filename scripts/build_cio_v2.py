#!/usr/bin/env python3
"""Build a separate CIO edition from saved facts. No model calls or original-file writes."""
from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGB
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from zeroops import paths
from zeroops import cio_v2_paths as loc
from zeroops.schema import incident_questions

NAVY, TEAL, INK, GREY, LIGHT = "112D44", "007F82", "233B4D", "586974", "EFF4F6"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def context(f):
    b, j = f["backends"]["jev_native"], f["jev"]
    strict = next(s for s in j["sensitivity"] if s["thresholds"].get("page_min") == 0.8)
    return dict(n=f["n"], templates=f["n_templates"], corpus=f["corpus_version"],
                questions=len(incident_questions()), p50=round(b["p50_ms"]),
                p95=round(b["p95_ms"]), cost=f'{b["cost_per_1k_usd"]:.3f}',
                exact=b["action_exact"]["k"], exact_pct=round(b["action_exact"]["pct"]),
                auto=b["auto"], auto_pct=round(b["auto_pct"]), auto_exact=b["auto_exact"],
                unsafe=b["unsafe_cases"], unsafe_auto=b["unsafe_auto"],
                unsafe_templates=b["unsafe_templates"], upper=round(b["unsafe_templates_upper95_pct"]),
                page=f["thresholds"]["page_min"], strict_page=strict["thresholds"]["page_min"],
                strict_auto=strict["auto"], rollback_proposed=j["first_option_proposed"],
                rollback_labelled=j["first_option_labelled"])


def fill(value, ctx):
    if isinstance(value, str):
        return value.format_map(ctx)
    if isinstance(value, list):
        return [fill(x, ctx) for x in value]
    if isinstance(value, dict):
        return {k: fill(v, ctx) for k, v in value.items()}
    return value


def box(slide, x, y, w, h, color, kind=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(color)
    s.line.fill.background()
    s._element.spPr.append(OxmlElement("a:effectLst"))
    # LibreOffice ignores the empty effectLst and draws the theme shadow from the shape style; drop the style.
    s._element.remove(s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style"))
    return s


def text(slide, value, x, y, w, h, size=18, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(3)
    return shape


def table(slide, headers, rows, x=.55, y=2.05, widths=None, row_h=.83, font=17):
    widths = widths or [3.0, 4.55, 4.7]
    offsets = [x]
    for w in widths:
        offsets.append(offsets[-1] + w)
    for c, header in enumerate(headers):
        box(slide, offsets[c], y, widths[c]-.025, .48, NAVY)
        text(slide, header, offsets[c]+.12, y+.08, widths[c]-.25, .34, 13, "FFFFFF", True)
    for r, row in enumerate(rows):
        ry = y+.51+r*row_h
        for c, value in enumerate(row):
            box(slide, offsets[c], ry, widths[c]-.025, row_h-.03, LIGHT if r % 2 == 0 else "F8FAFB")
            text(slide, str(value), offsets[c]+.12, ry+.12, widths[c]-.27, row_h-.17,
                 font, INK, c == 0)


def checks_rows(example):
    names = {"needs_page": "Page a person", "real_incident": "Real incident",
             "action_conf": "Action confidence", "action_risk": "Action risk",
             "severity": "Severity", "deploy_correlated": "Deploy correlation",
             "safe_to_autorollback": "Rollback score"}
    return [[names.get(c["check"], c["check"]),
             f'{c["value"]:.2f} {c["op"]} {c["threshold"]:.2f}',
             f'{c["margin"]:+.2f} / {"pass" if c["passed"] else "fail"}']
            for c in example["checks"]]


def build_deck(story, facts, target):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "ZeroOps: CIO decision brief — version 2"
    prs.core_properties.subject = "Editorial synthesis of the project's synthetic incident evaluation"
    prs.core_properties.author = ""
    for index, spec in enumerate(story["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box(slide, 0, 0, 13.333, .08, TEAL)
        text(slide, spec["section"].upper(), .55, .3, 12.2, .28, 11, TEAL, True)
        text(slide, spec["title"], .55, .78, 12.15, .94, 27, NAVY, True)
        text(slide, spec["subtitle"], .55, 1.72, 12.15, .36, 13, GREY)
        layout = spec["layout"]
        if layout in {"cards", "flow", "metrics"}:
            for c, (heading, body) in enumerate(spec["cards"]):
                x = .55 + c*4.15
                box(slide, x, 2.3, 3.98, 3.25, LIGHT)
                box(slide, x, 2.3, 3.98, .045, TEAL)
                text(slide, heading, x+.2, 2.55, 3.56, .82,
                     37 if layout == "metrics" else 20, TEAL if layout == "metrics" else NAVY, True)
                text(slide, body, x+.2, 3.43, 3.54, 1.94, 18)
        elif layout == "table":
            table(slide, spec["headers"], spec["rows"], y=2.3,
                  row_h=1.03 if len(spec["rows"]) == 3 else .80, font=16)
        elif layout == "comparison":
            rows = []
            for b in facts["backends"].values():
                rows.append([b["label"], f'{b["action_exact"]["k"]}/{b["n"]}',
                             f'{b["action_acceptable"]["k"]}/{b["n"]}',
                             f'{b["auto"]}/{b["n"]}',
                             f'{b["auto_exact"]}/{b["auto"]}' if b["auto"] else "—",
                             f'{b["p50_ms"]:.0f}/{b["p95_ms"]:.0f}',
                             f'${b["cost_per_1k_usd"]:.3f}' if b["cost_per_1k_usd"] else "GPU uncosted"])
            table(slide, ["Backend", "Exact action", "Acceptable*", "AUTO", "Exact / AUTO", "p50/p95 ms", "API $ / 1,000"],
                  rows, y=2.35, widths=[3.1, 1.43, 1.45, 1.05, 1.62, 1.62, 1.96], row_h=.65, font=14)
            text(slide, "*Acceptable-action criteria were defined after the first run. Use exact action match as the primary measure.",
                 .55, 5.6, 12.2, .46, 13, GREY)
        elif layout == "examples":
            for c, key in enumerate(("auto", "stopped")):
                e = facts["jev"]["examples"][key]
                x = .55+c*6.3
                text(slide, f'{e["id"]} · {e["archetype"].replace("_", " ")} · {e["outcome"]}',
                     x, 2.23, 6.0, .38, 17, NAVY, True)
                rows = checks_rows(e)
                table(slide, ["Check", "Value / rule", "Margin / result"], rows,
                      x=x, y=2.75, widths=[2.35, 1.7, 1.95], row_h=.36, font=11)
        box(slide, .55, 6.14, 12.23, .58, NAVY)
        text(slide, spec["takeaway"], .73, 6.26, 11.88, .35, 15, "FFFFFF", True)
        text(slide, spec["source"], .55, 6.92, 11.75, .40, 9, GREY)
        text(slide, str(index), 12.4, 6.96, .35, .22, 10, GREY)
        slide.notes_slide.notes_text_frame.text = spec["notes"] + "\n\n" + spec["source"]
    prs.save(target)


def add_runs(p, content):
    for i, part in enumerate(re.split(r"\*\*(.*?)\*\*", content)):
        p.add_run(part).bold = i % 2 == 1


def build_docx(memo, target):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = DocInches(8.27), DocInches(11.69)
    sec.top_margin = sec.bottom_margin = DocInches(.53)
    sec.left_margin = sec.right_margin = DocInches(.65)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", DocPt(10)
    normal.paragraph_format.space_after = DocPt(5)
    normal.paragraph_format.line_spacing = 1.02
    for name in ("Title", "Heading 2"):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.color.rgb = DocRGB.from_string(NAVY)
    doc.styles["Title"].font.size = DocPt(21)
    doc.styles["Title"].paragraph_format.space_after = DocPt(5)
    doc.styles["Heading 2"].font.size = DocPt(11)
    doc.styles["Heading 2"].paragraph_format.space_before = DocPt(7)
    doc.styles["Heading 2"].paragraph_format.space_after = DocPt(3)
    for line in memo.splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 2")
        elif line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = DocInches(.12)
            add_runs(p, "• " + line[2:])
        elif line.startswith("*") and not line.startswith("**"):
            p = doc.add_paragraph()
            r = p.add_run(line.strip("*"))
            r.italic, r.font.size = True, DocPt(8)
        else:
            add_runs(doc.add_paragraph(), line)
    doc.core_properties.title = "ZeroOps: CIO decision brief — version 2"
    doc.core_properties.author = ""
    doc.save(target)


def export_pdf(source, output):
    with tempfile.TemporaryDirectory(prefix="cio-v2-lo-") as profile:
        result = subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}",
                                 "--headless", "--convert-to", "pdf", "--outdir", str(output), str(source)],
                                capture_output=True, text=True, check=True, timeout=120)
        print(result.stdout.strip())
    pdf = output / (source.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError(f"PDF export missing: {pdf}")
    return pdf


def main():
    # Snapshot existing deliverables and inputs, excluding our separate edition.
    originals = [p for directory in (paths.CONTENT, paths.DELIVERABLES, paths.SCRIPTS)
                 for p in directory.iterdir() if p.is_file() and p.name != Path(__file__).name]
    before = {str(p): digest(p) for p in originals}
    facts = json.loads(paths.FACTS.read_text())
    ctx = context(facts)
    story = fill(json.loads(loc.STORY.read_text()), ctx)
    memo = loc.MEMO_TEMPLATE.read_text().format_map(ctx)
    output = loc.OUTPUT
    if output.exists():
        number = 2
        while (loc.OUTPUT / f"build_{number:02d}").exists():
            number += 1
        output = loc.OUTPUT / f"build_{number:02d}"
    output.mkdir(parents=True, exist_ok=False)
    deck, docx = output / loc.DECK.name, output / loc.MEMO_DOCX.name
    (output / loc.MEMO.name).write_text(memo)
    build_deck(story, facts, deck)
    build_docx(memo, docx)
    (output / "editorial-notes.md").write_text(loc.EDITORIAL.read_text())
    # Preserve a readable, searchable version of every slide, including speaker notes.
    script = []
    for i, s in enumerate(story["slides"], 1):
        script.extend([f'## {i}. {s["title"]}', s["subtitle"], ""])
        if "cards" in s:
            script.extend(f"**{h}**\n\n{b}\n" for h, b in s["cards"])
        if "rows" in s:
            script.extend(" | ".join(row) for row in [s["headers"], *s["rows"]])
        if s["layout"] == "comparison":
            for b in facts["backends"].values():
                script.append(f'{b["label"]}: exact {b["action_exact"]["k"]}/{b["n"]}; acceptable {b["action_acceptable"]["k"]}/{b["n"]}; AUTO {b["auto"]}/{b["n"]}; exact within AUTO {b["auto_exact"]}/{b["auto"]}.')
        if s["layout"] == "examples":
            for e in facts["jev"]["examples"].values():
                script.append(f'{e["id"]}: {e["outcome"]}\n' + "\n".join(" | ".join(row) for row in checks_rows(e)))
        script.extend(["", s["takeaway"], "", s["source"], "", "Speaker notes: " + s["notes"], ""])
    (output / "slide-narrative.md").write_text("\n".join(script))
    deck_pdf, memo_pdf = export_pdf(deck, output), export_pdf(docx, output)
    previews = output / loc.PREVIEWS.name
    previews.mkdir()
    subprocess.run(["pdftoppm", "-scale-to", "1400", "-png", str(deck_pdf), str(previews / "slide")],
                   check=True, timeout=120)
    subprocess.run(["pdftoppm", "-scale-to", "1600", "-png", str(memo_pdf), str(previews / "memo")],
                   check=True, timeout=120)
    links = [(deck_pdf.name, "Presentation · PDF"), (deck.name, "Presentation · editable PowerPoint"),
             (memo_pdf.name, "Executive memo · PDF"), (docx.name, "Executive memo · editable Word"),
             (loc.MEMO.name, "Executive memo · Markdown"), ("slide-narrative.md", "Slide text and speaker notes"),
             ("editorial-notes.md", "Editorial review")]
    (output / "README.md").write_text("# ZeroOps — CIO edition, version 2\n\n" +
        "A separate revision of the deck and executive memo. Ten core slides and two appendices.\n\n" +
        "\n".join(f"- [{label}]({name})" for name, label in links) +
        "\n\nMeasured figures come from the saved project results. See editorial notes for scope and rebuild instructions.\n")
    (output / "index.html").write_text('<!doctype html><html lang="en"><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1"><title>ZeroOps — CIO edition</title>'
        '<style>body{font:18px/1.6 system-ui;max-width:850px;margin:70px auto;padding:0 24px;color:#112d44}'
        'a{color:#007f82}li{margin:12px 0}h1{line-height:1.15}</style>'
        '<body><p>CIO DECISION BRIEF · VERSION 2</p><h1>Test routine incident decisions before expanding automation</h1>'
        '<p>Ten core slides, two evidence appendices and a companion executive memo.</p><ul>' +
        ''.join(f'<li><a href="{html.escape(n)}">{html.escape(label)}</a></li>' for n, label in links) + '</ul></body></html>')
    changed = [p for p, value in before.items() if digest(Path(p)) != value]
    if changed:
        raise RuntimeError(f"Original files changed: {changed}")
    (output / "build-manifest.json").write_text(json.dumps({
        "corpus_version": facts["corpus_version"], "facts_sha256": digest(paths.FACTS),
        "source_sha256": {p.name: digest(p) for p in (loc.STORY, loc.MEMO_TEMPLATE, loc.EDITORIAL)},
        "original_files_verified_unchanged": len(before), "slide_count": len(story["slides"]),
        "artifacts_sha256": {p.name: digest(p) for p in output.iterdir() if p.is_file()}
    }, indent=2))
    print(f"Created {output}; {len(before)} original files verified unchanged.")


if __name__ == "__main__":
    main()
